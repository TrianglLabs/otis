#!/usr/bin/env python3
"""Local document generation for Otis. No network calls or automatic installs."""

import argparse
import importlib.util
import io
import json
import math
import os
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from contextlib import closing
from pathlib import Path
from xml.etree import ElementTree
from xml.sax.saxutils import escape

MAX_BYTES = 20 * 1024 * 1024
PACKAGES = {
    "docx": "python-docx",
    "reportlab": "reportlab",
    "pypdf": "pypdf",
    "pypdfium2": "pypdfium2",
    "PIL": "pillow",
}


def require(*modules):
    missing = [PACKAGES[name] for name in modules if importlib.util.find_spec(name) is None]
    if missing:
        raise ValueError(
            "Missing local dependencies: "
            + ", ".join(missing)
            + ". Run this operation through Otis's document tool to prepare its private dependencies."
        )


def office_command():
    bundled = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    found = shutil.which("soffice") or shutil.which("libreoffice")
    return found or (str(bundled) if bundled.is_file() else None)


def check(_args):
    available = {name: importlib.util.find_spec(name) is not None for name in PACKAGES}
    office = office_command()
    return {
        "python": sys.version.split()[0],
        "missing_packages": [package for name, package in PACKAGES.items() if not available[name]],
        "libreoffice": bool(office),
        "capabilities": {
            "create_docx": available["docx"],
            "create_pdf": available["reportlab"] and available["pypdf"],
            "convert_docx_to_pdf": bool(office) and available["docx"] and available["pypdf"],
            "render_pdf": available["pypdfium2"] and available["PIL"],
            "inspect_pdf": available["pypdfium2"] and available["pypdf"],
            "edit_pdf": available["pypdfium2"] and available["pypdf"] and available["PIL"],
        },
    }


def workspace_path(value, *, output=False):
    path = Path(value).resolve()
    if not path.is_relative_to(Path.cwd().resolve()):
        raise ValueError("Document inputs and outputs must be inside the workspace.")
    if output:
        if path.exists() or Path(value).is_symlink():
            raise ValueError("Output already exists; choose a new path.")
        if not path.parent.is_dir():
            raise ValueError("Output parent directory does not exist.")
    elif not path.is_file() or not 0 < path.stat().st_size <= MAX_BYTES:
        raise ValueError("Input must be a nonempty regular file of at most 20 MB.")
    return path


def keys(value, allowed, label):
    if not isinstance(value, dict):
        raise ValueError(label + " must be an object.")
    unknown = set(value) - set(allowed)
    if unknown:
        raise ValueError(label + " has unknown options: " + ", ".join(sorted(unknown)))


def number(value, minimum, maximum, label, *, integer=False):
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(value)
        or not minimum <= value <= maximum
        or (integer and not isinstance(value, int))
    ):
        raise ValueError(
            f"{label} must be {'an integer' if integer else 'a number'} from {minimum} to {maximum}."
        )
    return value


def text(value):
    if not isinstance(value, str) or not value.strip():
        raise ValueError("Document text must be a nonempty string.")
    if any(
        (ord(c) < 32 and c != "\n") or 0xD800 <= ord(c) <= 0xDFFF or c in "\ufffe\uffff"
        for c in value
    ):
        raise ValueError(
            "Document text contains unsupported control characters, tabs, or invalid Unicode."
        )
    return value


def normalized(value):
    return "".join(unicodedata.normalize("NFC", value).split())


def verify_text(actual, expected):
    remaining = normalized(actual)
    for value in expected:
        needle = normalized(value)
        index = remaining.find(needle)
        if index < 0:
            raise ValueError(
                "Output verification failed: expected text is missing or out of order: "
                + value[:100]
            )
        remaining = remaining[index + len(needle) :]


def pdf_text(data, max_pages=500):
    require("pypdf")
    from pypdf import PdfReader

    pdf = PdfReader(io.BytesIO(data), strict=True)
    if pdf.is_encrypted:
        raise ValueError("Encrypted PDFs are not supported by this helper.")
    if not 1 <= len(pdf.pages) <= max_pages:
        raise ValueError(f"PDF has {len(pdf.pages)} pages; the allowed maximum is {max_pages}.")
    return "\n".join(page.extract_text() or "" for page in pdf.pages), len(pdf.pages)


def docx_text(data):
    require("docx")
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    # Match the bounded input contract of Otis's native document tools. Linked assets are not fetched.
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(entry.file_size for entry in entries) > 128 * 1024 * 1024:
            raise ValueError("Word archive exceeds the expanded document limit.")
        for entry in entries:
            if entry.filename == "word/document.xml" and entry.file_size > 16 * 1024 * 1024:
                raise ValueError("Word main document XML exceeds 16 MB.")
            if "vbaproject" in entry.filename.lower():
                raise ValueError("Macro-enabled Word packages are not supported.")
            if not entry.filename.endswith(".rels"):
                continue
            if entry.file_size > 2 * 1024 * 1024:
                raise ValueError("Word relationships exceed the document limit.")
            if any(
                relation.get("TargetMode") == "External"
                and not relation.get("Type", "").endswith("/hyperlink")
                for relation in ElementTree.fromstring(archive.read(entry))
            ):
                raise ValueError(
                    "Word contains linked external assets; embed them before conversion."
                )
    values = []
    for item in Document(io.BytesIO(data)).iter_inner_content():
        if isinstance(item, Paragraph) and item.text.strip():
            values.append(item.text)
        elif isinstance(item, Table):
            values.extend(cell.text for row in item.rows for cell in row.cells if cell.text.strip())
    return values


def create_docx(spec):
    require("docx")
    from docx import Document
    from docx.shared import Mm, Pt

    document = Document()
    document.core_properties.title = spec.get("title", "")
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    section = document.sections[0]
    width, height = (210, 297) if spec.get("page_size") == "a4" else (215.9, 279.4)
    section.page_width, section.page_height = Mm(width), Mm(height)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(
        spec.get("margin_mm", 18)
    )
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(spec.get("font_size", 11))
    normal.paragraph_format.space_after = Pt(6)
    for block in spec["blocks"]:
        kind = block["type"]
        if kind == "heading":
            document.add_heading(block["text"], level=block.get("level", 1))
        elif kind == "paragraph":
            document.add_paragraph(block["text"])
        elif kind == "bullets":
            for value in block["items"]:
                document.add_paragraph(value, style="List Bullet")
        elif kind == "table":
            table = document.add_table(rows=0, cols=len(block["rows"][0]))
            table.style = "Table Grid"
            for values in block["rows"]:
                for cell, value in zip(table.add_row().cells, values):
                    cell.text = value
        else:
            document.add_page_break()
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def create_pdf(spec, expected):
    require("reportlab", "pypdf")
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    characters = set("".join(expected) + "•") - {"\n"}
    if any(unicodedata.bidirectional(c) in ("R", "AL", "AN") for c in characters):
        raise ValueError(
            "This PDF helper does not shape right-to-left text. Use a suitable document renderer."
        )
    candidates = (
        [spec["font_path"]]
        if "font_path" in spec
        else [
            "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
            str(Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "arial.ttf"),
            str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"),
        ]
    )
    for path in candidates:
        if not Path(path).is_file():
            continue
        font = TTFont("OtisDocument", path)
        if all(font.face.charToGlyph.get(ord(c), 0) != 0 for c in characters):
            pdfmetrics.registerFont(font)
            break
    else:
        raise ValueError(
            "No available TrueType font covers the document text. Set font_path to a suitable local font."
        )
    stream = io.BytesIO()
    margin = spec.get("margin_mm", 18) * mm
    size = spec.get("font_size", 11)
    page = A4 if spec.get("page_size") == "a4" else letter
    document = SimpleDocTemplate(
        stream,
        pagesize=page,
        topMargin=margin,
        bottomMargin=margin,
        leftMargin=margin,
        rightMargin=margin,
        title=spec.get("title", ""),
        author="",
    )
    body = ParagraphStyle(
        "Body",
        fontName="OtisDocument",
        fontSize=size,
        leading=size * 1.35,
        spaceAfter=6,
        allowWidows=0,
        allowOrphans=0,
    )
    bullet = ParagraphStyle(
        "Bullet",
        parent=body,
        leftIndent=12,
        bulletIndent=0,
        bulletFontName="OtisDocument",
        bulletFontSize=size,
    )

    def paragraph(value, style=body, **kwargs):
        return Paragraph(escape(value).replace("\n", "<br/>"), style, **kwargs)

    story = []
    for block in spec["blocks"]:
        kind = block["type"]
        if kind == "heading":
            scale = [2, 1.4, 1.2, 1.1][block.get("level", 1)]
            style = ParagraphStyle(
                "Heading",
                parent=body,
                fontSize=size * scale,
                leading=size * scale * 1.2,
                spaceBefore=8,
                spaceAfter=6,
                keepWithNext=True,
            )
            story.append(paragraph(block["text"], style))
        elif kind == "paragraph":
            story.append(paragraph(block["text"]))
        elif kind == "bullets":
            story.extend(paragraph(value, bullet, bulletText="•") for value in block["items"])
        elif kind == "table":
            rows = [[paragraph(cell) for cell in row] for row in block["rows"]]
            table = Table(
                rows,
                colWidths=[(page[0] - 2 * margin - 12) / len(rows[0])] * len(rows[0]),
                repeatRows=1,
            )
            table.setStyle(
                TableStyle(
                    [
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
                    ]
                )
            )
            story.extend([table, Spacer(1, 8)])
        else:
            story.append(PageBreak())
    document.build(story)
    return stream.getvalue()


def publish(data, output):
    if not 0 < len(data) <= MAX_BYTES:
        raise ValueError("Generated output is empty or exceeds 20 MB.")
    descriptor, temporary = tempfile.mkstemp(prefix=".otis-document-", dir=output.parent)
    try:
        with os.fdopen(descriptor, "wb") as file:
            file.write(data)
        os.link(temporary, output)  # Atomic no-clobber publication on the same filesystem.
    finally:
        os.unlink(temporary)


def create(args):
    output = workspace_path(args.output, output=True)
    extension = output.suffix.lower()
    if extension not in (".pdf", ".docx"):
        raise ValueError("Output must be .pdf or .docx; no format fallback is performed.")
    path = workspace_path(args.spec)
    if path.stat().st_size > 1024 * 1024:
        raise ValueError("Specification exceeds 1 MB.")
    spec = json.loads(path.read_text(encoding="utf-8"))
    keys(
        spec,
        ["title", "page_size", "margin_mm", "font_size", "font_path", "max_pages", "blocks"],
        "Specification",
    )
    blocks = spec.get("blocks")
    if not isinstance(blocks, list) or not 1 <= len(blocks) <= 1000:
        raise ValueError("blocks must contain 1–1000 entries.")
    if spec.get("page_size", "letter") not in ("letter", "a4"):
        raise ValueError("page_size must be letter or a4.")
    number(spec.get("margin_mm", 18), 10, 40, "margin_mm")
    number(spec.get("font_size", 11), 9, 16, "font_size")
    if "title" in spec:
        text(spec["title"])
    if "font_path" in spec:
        text(spec["font_path"])
        if extension != ".pdf":
            raise ValueError("font_path is PDF-only; DOCX uses Arial.")
    if "max_pages" in spec:
        number(spec["max_pages"], 1, 500, "max_pages", integer=True)
        if extension != ".pdf":
            raise ValueError("max_pages is PDF-only. Convert DOCX to PDF to check its pagination.")
    expected = []
    previous_break = True
    for block in blocks:
        if not isinstance(block, dict):
            raise ValueError("Each block must be an object.")
        kind = block.get("type")
        if kind in ("heading", "paragraph"):
            keys(
                block, ["type", "text", "level"] if kind == "heading" else ["type", "text"], "Block"
            )
            expected.append(text(block.get("text")))
            if kind == "heading":
                number(block.get("level", 1), 0, 3, "heading level", integer=True)
        elif kind == "bullets":
            keys(block, ["type", "items"], "Bullet block")
            items = block.get("items")
            if not isinstance(items, list) or not 1 <= len(items) <= 200:
                raise ValueError("Bullet items must contain 1–200 strings.")
            expected.extend(text(item) for item in items)
        elif kind == "table":
            keys(block, ["type", "rows"], "Table block")
            rows = block.get("rows")
            if not isinstance(rows, list) or not 1 <= len(rows) <= 200:
                raise ValueError("Tables must contain 1–200 rows.")
            columns = len(rows[0]) if isinstance(rows[0], list) else 0
            if not 1 <= columns <= 8:
                raise ValueError("Tables must contain 1–8 columns.")
            for row in rows:
                if not isinstance(row, list) or len(row) != columns:
                    raise ValueError("Table rows must have equal column counts.")
                expected.extend(text(cell) for cell in row)
        elif kind == "page_break":
            keys(block, ["type"], "Page break")
            if previous_break:
                raise ValueError("Page breaks cannot be first or adjacent.")
        else:
            raise ValueError("Unsupported block type: " + str(kind))
        previous_break = kind == "page_break"
    if previous_break:
        raise ValueError("Page breaks cannot be last.")
    if sum(map(len, expected)) > 200_000:
        raise ValueError("Document text exceeds 200,000 characters.")
    if extension == ".pdf":
        data = create_pdf(spec, expected)
        actual, pages = pdf_text(data, spec.get("max_pages", 500))
    else:
        data = create_docx(spec)
        actual, pages = "\n".join(docx_text(data)), None
    verify_text(actual, expected)
    publish(data, output)
    return {
        "path": str(output),
        "format": extension[1:],
        "pages": pages,
        "verified": ["reopened", "expected_text"],
        "visual_review": "not_performed",
    }


def convert(args):
    source = workspace_path(args.source)
    output = workspace_path(args.output, output=True)
    if source.suffix.lower() != ".docx" or output.suffix.lower() != ".pdf":
        raise ValueError("Conversion supports DOCX to PDF only.")
    require("docx", "pypdf")
    office = office_command()
    if not office:
        raise ValueError(
            "LibreOffice is unavailable. DOCX-to-PDF conversion requires a local LibreOffice installation."
        )
    source_bytes = source.read_bytes()
    expected = docx_text(source_bytes)
    if not expected:
        raise ValueError(
            "Source has no body/table text to verify; conversion needs another review workflow."
        )
    with tempfile.TemporaryDirectory(prefix="otis-convert-") as directory:
        staging = Path(directory)
        (staging / "source.docx").write_bytes(source_bytes)
        subprocess.run(
            [
                office,
                "-env:UserInstallation=" + (staging / "profile").as_uri(),
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(staging),
                str(staging / "source.docx"),
            ],
            check=True,
            capture_output=True,
            timeout=90,
        )
        result = staging / "source.pdf"
        if not result.is_file() or result.stat().st_size > MAX_BYTES:
            raise ValueError("LibreOffice did not produce a bounded PDF output.")
        data = result.read_bytes()
    actual, pages = pdf_text(data)
    verify_text(actual, expected)
    publish(data, output)
    return {
        "path": str(output),
        "format": "pdf",
        "pages": pages,
        "verified": ["reopened", "body_and_table_text"],
        "visual_review": "not_performed",
    }


def render(args):
    source = workspace_path(args.source)
    output = workspace_path(args.output_dir, output=True)
    if source.suffix.lower() != ".pdf":
        raise ValueError("Render a PDF; convert Word documents first.")
    require("pypdfium2", "PIL")
    import pypdfium2

    with pypdfium2.PdfDocument(source) as pdf:
        pages = (
            [int(value) for value in args.pages.split(",")]
            if args.pages
            else list(range(1, len(pdf) + 1))
        )
        if (
            not 1 <= len(pages) <= 20
            or len(set(pages)) != len(pages)
            or any(p < 1 or p > len(pdf) for p in pages)
        ):
            raise ValueError("Render 1–20 unique pages; use --pages 1,2,3 to select a subset.")
        pdf.init_forms()
        with tempfile.TemporaryDirectory(prefix="otis-render-", dir=output.parent) as directory:
            staging = Path(directory)
            for page in pages:
                image = staging / f"page-{page}.png"
                with closing(pdf[page - 1]) as source_page:
                    width, height = source_page.get_size()
                    if (
                        not math.isfinite(width * height)
                        or min(width, height) <= 0
                        or width * height * (120 / 72) ** 2 > 20_000_000
                    ):
                        raise ValueError("Page dimensions exceed the render limit.")
                    with closing(source_page.render(scale=120 / 72)) as bitmap:
                        bitmap.to_pil().save(image)
                if not 0 < image.stat().st_size <= MAX_BYTES:
                    raise ValueError("Rendered page exceeds the image size limit.")
            output.mkdir(mode=0o700)  # Fail if another process created it meanwhile.
            try:
                for page in pages:
                    target = output / f"page-{page}.png"
                    shutil.copyfile(staging / target.name, target)
                    target.chmod(0o600)
            except Exception:
                shutil.rmtree(output)
                raise
    return {
        "directory": str(output),
        "images": [f"page-{p}.png" for p in pages],
        "visual_review": "not_performed",
    }


def inspect_pdf(args):
    require("pypdfium2", "pypdf")
    from pdf_edit import inspect

    source = workspace_path(args.source)
    if source.suffix.lower() != ".pdf":
        raise ValueError("PDF inspection requires a .pdf source.")
    pages = [int(value) for value in args.pages.split(",")] if args.pages else None
    return inspect(source.read_bytes(), pages)


def edit_pdf(args):
    require("pypdfium2", "pypdf", "PIL")
    from pdf_edit import edit

    source = workspace_path(args.source)
    output = workspace_path(args.output, output=True)
    plan = workspace_path(args.spec)
    if source.suffix.lower() != ".pdf" or output.suffix.lower() != ".pdf":
        raise ValueError("PDF editing requires .pdf source and output paths.")
    if plan.stat().st_size > 1024 * 1024:
        raise ValueError("PDF edit plan exceeds 1 MB.")
    data, verification = edit(source.read_bytes(), json.loads(plan.read_text(encoding="utf-8")))
    publish(data, output)
    return {"path": str(output), "format": "pdf", **verification}


def main():
    commands = {
        "check": (check, []),
        "create": (create, ["--spec", "--output"]),
        "convert": (convert, ["--source", "--output"]),
        "render": (render, ["--source", "--output-dir", "--pages"]),
        "inspect-pdf": (inspect_pdf, ["--source", "--pages"]),
        "edit-pdf": (edit_pdf, ["--source", "--spec", "--output"]),
    }
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    for name, (_, flags) in commands.items():
        subparser = subparsers.add_parser(name)
        for flag in flags:
            subparser.add_argument(flag, required=flag != "--pages")
    args = parser.parse_args()
    try:
        if sys.version_info < (3, 10):
            raise ValueError("Python 3.10 or later is required.")
        result = commands[args.command][0](args)
        print(json.dumps({"ok": True, **result}, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({"ok": False, "error": str(error)}, ensure_ascii=False), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
